#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
实习手册封面生成器（GUI / 打包版）
=====================================================================
功能
  1. 输入班级名 + 份数，实时预览 A4 封面（预览即打印效果）
  2. 批量列表：一次做好几个班的封面，逐个或一次全打
  3. 下拉选择打印机，直接打印，可设置打印份数
  4. 可导出 PNG / Word(.docx)

打印原理
  程序自己按打印机原生 DPI 画 A4 位图，再送进打印机 DC。
  所以「预览 = 打印结果」，不装 Word 也能打。

打包成 exe（Windows）
  双击 build.bat
"""

import os
import sys
import json
import platform
import threading
import queue

from PIL import Image, ImageDraw, ImageFont

import tkinter as tk
from tkinter import ttk, messagebox, filedialog

# ============================ 常量 ============================

IS_WIN = platform.system() == "Windows"

A4_W_MM = 210.0
A4_H_MM = 297.0

PRINT_DPI_MAX = 600        # 打印位图最高分辨率（跟随打印机原生 DPI）
PREVIEW_DPI = 200          # 预览渲染分辨率

DEFAULT_FONT_SIZE = 36
DEFAULT_SUBTITLE = "实习手册"
DEFAULT_COPIES = 1

# 排版几何（单位 pt，1pt = 1/72 英寸）
PT_PER_MM = 72.0 / 25.4
PAGE_W_PT = A4_W_MM * PT_PER_MM
PAGE_H_PT = A4_H_MM * PT_PER_MM
MARGIN_PT = 1.5 * 10 * PT_PER_MM     # 1.5cm
BLOCK_CENTER_RATIO = 0.44            # 竖排文字块中心所在页高比例
FOOT_GAP_RATIO = 0.9                 # 「XX份」离下边距的距离（相对其字号）

CONFIG_NAME = "cover_maker_config.json"


# ============================ 字体 ============================

def find_cjk_font():
    """按优先级找一款中文黑体。找不到返回 None。"""
    cands = []

    if IS_WIN:
        windir = os.environ.get("WINDIR", r"C:\Windows")
        fonts = windir + r"\Fonts"
        cands += [
            os.path.join(fonts, "simhei.ttf"),    # 黑体
            os.path.join(fonts, "msyh.ttc"),      # 微软雅黑
            os.path.join(fonts, "msyhl.ttc"),
            os.path.join(fonts, "simsun.ttc"),    # 宋体
            os.path.join(fonts, "simkai.ttf"),    # 楷体
            os.path.join(fonts, "simfang.ttf"),   # 仿宋
        ]
    else:
        cands += [
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/opentype/noto/NotoSerifCJK-Bold.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
            "/System/Library/Fonts/PingFang.ttc",
        ]

    base = os.path.dirname(os.path.abspath(sys.argv[0]))
    cands = [os.path.join(base, "simhei.ttf"), os.path.join(base, "font.ttf")] + cands

    for p in cands:
        if os.path.exists(p):
            try:
                ImageFont.truetype(p, 32)
                return p
            except Exception:
                continue
    return None


FONT_PATH = find_cjk_font()
_FONT_CACHE = {}


def load_font(pt, dpi):
    key = (round(pt, 1), dpi)
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    size = max(1, int(round(pt * dpi / 72.0)))
    if FONT_PATH:
        font = ImageFont.truetype(FONT_PATH, size)
    else:
        font = ImageFont.load_default()
    _FONT_CACHE[key] = font
    return font


# ============================ 文本 ============================

def split_vertical(text: str):
    """班级名拆成竖排行：连续数字/字母归一组，汉字单字一行。
    例：'24药剂01' -> ['24', '药', '剂', '01']
    """
    parts, i = [], 0
    text = (text or "").strip()
    while i < len(text):
        ch = text[i]
        if ch.isascii() and ch.isalnum():
            j = i
            while j < len(text) and text[j].isascii() and text[j].isalnum():
                j += 1
            parts.append(text[i:j])
            i = j
        else:
            parts.append(ch)
            i += 1
    return parts


# ============================ 预览图片转换 ============================

def make_photo(img):
    """把 PIL 图片转成 tkinter 可显示的 PhotoImage。

    刻意不走 PIL.ImageTk：Pillow 10+ 的懒加载机制会让 PyInstaller 把它
    当成「可选模块」漏打包，导出的 exe 里预览会是一片空白。
    这里改用 tkinter 原生 PhotoImage（PNG 优先，GIF 兜底）。
    """
    import io
    import base64

    for fmt in ("PNG", "GIF"):
        try:
            buf = io.BytesIO()
            img.save(buf, format=fmt)
            data = base64.b64encode(buf.getvalue()).decode("ascii")
            return tk.PhotoImage(data=data)
        except Exception:
            continue

    # 最后兜底：PIL 自带（源码运行时的正常路径；打包后可能不可用）
    try:
        from PIL import ImageTk
        return ImageTk.PhotoImage(img)
    except Exception as ex:
        raise RuntimeError(f"无法生成预览图片：{ex}") from ex


# ============================ 排版几何 ============================

class Layout:
    """竖排封面的几何计算结果，单位 pt（相对页面左上角）。"""

    def __init__(self, class_name, copies, subtitle=DEFAULT_SUBTITLE,
                 font_size=DEFAULT_FONT_SIZE, voffset=0.0):
        self.class_name = (class_name or "").strip()
        self.copies = (copies or "").strip()
        self.subtitle = subtitle or ""
        self.font_size = font_size
        self.voffset = voffset

        self.lines = split_vertical(self.class_name) + \
                     [c for c in self.subtitle if c.strip()]

        self.main_pt = font_size
        self.foot_pt = font_size * 0.78
        self.line_pitch = font_size * 1.45

        self.block_top = BLOCK_CENTER_RATIO * PAGE_H_PT + \
                         (voffset / 100.0) * PAGE_H_PT - \
                         len(self.lines) * self.line_pitch / 2.0
        self.block_height = len(self.lines) * self.line_pitch
        self.foot_line = self.foot_pt * 1.2
        # 底部「XX份」的中心线位置（贴着下边距上方）
        self.foot_center = PAGE_H_PT - MARGIN_PT - self.foot_pt * FOOT_GAP_RATIO

    def overflow(self):
        """内容是否超出页面（用于界面提示）。"""
        bottom = self.foot_center + self.foot_line / 2.0
        return self.block_top < MARGIN_PT or bottom > PAGE_H_PT - MARGIN_PT * 0.2

    # ---- 位图渲染 ----
    def render(self, dpi=300):
        W = int(round(A4_W_MM / 25.4 * dpi))
        H = int(round(A4_H_MM / 25.4 * dpi))
        img = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(img)

        main_font = load_font(self.main_pt, dpi)
        foot_font = load_font(self.foot_pt, dpi)

        cx = W / 2.0
        pitch_px = self.line_pitch / 72.0 * dpi
        for idx, line in enumerate(self.lines):
            y = self.block_top / 72.0 * dpi + pitch_px * (idx + 0.5)
            d.text((cx, y), line, font=main_font, fill="black", anchor="mm")

        fy = self.foot_center / 72.0 * dpi
        d.text((cx, fy), f"{self.copies}份", font=foot_font, fill="black", anchor="mm")

        return img


# ============================ 打印机 ============================

def list_printers():
    if not IS_WIN:
        return []
    try:
        import win32print
        flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
        return [p[2] for p in win32print.EnumPrinters(flags)]
    except Exception:
        return []


def get_default_printer():
    if not IS_WIN:
        return ""
    try:
        import win32print
        return win32print.GetDefaultPrinter()
    except Exception:
        return ""


def _set_a4_devmode(printer_name):
    """尽量把打印机纸张设为 A4 纵向。失败不致命。"""
    if not IS_WIN:
        return
    try:
        import win32print
        import win32con
        h = win32print.OpenPrinter(printer_name)
        try:
            props = win32print.GetPrinter(h, 2)
            dm = props.get("pDevMode")
            if dm is not None:
                dm.PaperSize = 9          # DMPAPER_A4
                dm.Orientation = 1        # 纵向
                try:
                    dm.Fields = dm.Fields | win32con.DM_PAPERSIZE | win32con.DM_ORIENTATION
                except Exception:
                    pass
                win32print.SetPrinter(h, 2, props, 0)
        finally:
            win32print.ClosePrinter(h)
    except Exception:
        pass


def print_layout(layout: Layout, printer_name: str, copies: int = 1):
    """把封面按整页送到指定打印机。"""
    if not IS_WIN:
        raise RuntimeError("只有 Windows 系统才能调用打印机。")

    import win32print
    import win32ui
    import win32con
    from PIL import ImageWin

    if not printer_name:
        printer_name = win32print.GetDefaultPrinter()
    if not printer_name:
        raise RuntimeError("没有可用打印机，请先安装或连接一台打印机。")

    _set_a4_devmode(printer_name)

    hdc = win32ui.CreateDC()
    hdc.CreatePrinterDC(printer_name)

    try:
        # 打印机原生 DPI，用它渲染才能得到清晰文字
        try:
            native_dpi = hdc.GetDeviceCaps(win32con.LOGPIXELSX)
        except Exception:
            native_dpi = 300
        dpi = max(200, min(int(native_dpi or 300), PRINT_DPI_MAX))

        pw = hdc.GetDeviceCaps(win32con.HORZRES)
        ph = hdc.GetDeviceCaps(win32con.VERTRES)

        img = layout.render(dpi=dpi)
        if img.mode != "RGB":
            img = img.convert("RGB")

        # 等比放进可打印区域，居中，避免拉伸变形
        scale = min(pw / img.width, ph / img.height)
        w = int(img.width * scale)
        h = int(img.height * scale)
        x = (pw - w) // 2
        y = (ph - h) // 2

        dib = ImageWin.Dib(img)

        hdc.StartDoc("实习手册封面")
        for _ in range(max(1, copies)):
            hdc.StartPage()
            dib.draw(hdc.GetHandleOutput(), (x, y, x + w, y + h))
            hdc.EndPage()
        hdc.EndDoc()
    finally:
        try:
            hdc.DeleteDC()
        except Exception:
            pass


def print_via_word(layout: Layout, printer_name: str, copies: int, docx_path: str):
    """用 Microsoft Word 打印（矢量文字，最清晰）。未装 Word 会抛异常。"""
    if not IS_WIN:
        raise RuntimeError("只有 Windows 系统才能调用打印机。")
    import win32com.client as win32
    word = win32.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(docx_path, ReadOnly=False)
        try:
            if printer_name:
                doc.ActivePrinter = printer_name
            doc.PrintOut(Copies=max(1, copies), PrintToFile=False)
        finally:
            doc.Close(SaveChanges=0)
    finally:
        try:
            word.Quit()
        except Exception:
            pass


# ============================ 导出 Word ============================

def export_docx(layout: Layout, path: str):
    from docx import Document
    from docx.shared import Pt, Cm, Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def set_run_font(run, name="黑体", size=36):
        run.font.name = name
        run.font.size = Pt(size)
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            rfonts.set(qn(attr), name)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(A4_W_MM)
    sec.page_height = Mm(A4_H_MM)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)

    for idx, line in enumerate(layout.lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        # 首行用段前距把整块文字顶到与位图预览相同的位置
        pf.space_before = Pt(max(0.0, layout.block_top - MARGIN_PT)) if idx == 0 else Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(layout.line_pitch)   # 固定行距，保证行位精确
        set_run_font(p.add_run(line), "黑体", int(round(layout.main_pt)))

    # 底部「XX份」：算好与文字块之间的空隙，落位同样与预览一致
    block_bottom = layout.block_top + layout.block_height
    gap = (layout.foot_center - layout.foot_line / 2.0) - block_bottom
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(max(0.0, gap))
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(layout.foot_line)
    set_run_font(p.add_run(f"{layout.copies}份"), "黑体", int(round(layout.foot_pt)))

    doc.save(path)
    return path


# ============================ 配置 ============================

def config_path():
    base = os.path.dirname(os.path.abspath(sys.argv[0]))
    if IS_WIN:
        try:
            base = os.path.join(os.environ.get("APPDATA", base), "CoverMaker")
            os.makedirs(base, exist_ok=True)
        except Exception:
            pass
    return os.path.join(base, CONFIG_NAME)


def load_config():
    try:
        with open(config_path(), "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_config(data):
    try:
        with open(config_path(), "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


# ============================ GUI ============================

class App(tk.Tk):

    def __init__(self):
        super().__init__()
        self.title("实习手册封面生成器")
        self.geometry("1040x780")
        self.minsize(940, 660)

        cfg = load_config()

        self.var_class = tk.StringVar(value=cfg.get("class_name", "24药剂01"))
        self.var_copies = tk.StringVar(value=str(cfg.get("copies", "47")))
        self.var_subtitle = tk.StringVar(value=cfg.get("subtitle", DEFAULT_SUBTITLE))
        self.var_fontsize = tk.IntVar(value=int(cfg.get("font_size", DEFAULT_FONT_SIZE)))
        self.var_voffset = tk.DoubleVar(value=float(cfg.get("voffset", 0)))
        self.var_printer = tk.StringVar(value=cfg.get("printer", ""))
        self.var_pcopies = tk.StringVar(value=str(cfg.get("print_copies", DEFAULT_COPIES)))
        self.var_word_print = tk.BooleanVar(value=bool(cfg.get("word_print", False)))

        self._photo = None
        self._rows = []            # [(班级, 份数)]
        self._q = queue.Queue()     # 工作线程 -> 主线程 的消息通道

        self._build_ui()
        self._refresh_printers()
        self._load_rows(cfg.get("rows", []))

        for v in (self.var_class, self.var_copies, self.var_subtitle,
                  self.var_fontsize, self.var_voffset):
            v.trace_add("write", lambda *_: self.schedule_preview())

        self.schedule_preview()
        self.after(80, self._poll_queue)
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---------------- 主线程消息轮询 ----------------
    def _poll_queue(self):
        try:
            while True:
                kind, payload = self._q.get_nowait()
                if kind == "status":
                    self.lbl_status.configure(text=payload)
                elif kind == "done":
                    self._done(payload[0], payload[1])
        except queue.Empty:
            pass
        self.after(80, self._poll_queue)

    # ---------------- 界面 ----------------
    def _build_ui(self):
        left = ttk.Frame(self, padding=(12, 10))
        left.pack(side="left", fill="y")
        left.configure(width=320)
        left.pack_propagate(False)

        right = ttk.Frame(self, padding=(4, 10, 12, 10))
        right.pack(side="left", fill="both", expand=True)

        # ---- 内容设置 ----
        box = ttk.LabelFrame(left, text="封面内容", padding=10)
        box.pack(fill="x")

        ttk.Label(box, text="班级名称：").grid(row=0, column=0, sticky="w", pady=3)
        e = ttk.Entry(box, textvariable=self.var_class, width=18)
        e.grid(row=0, column=1, sticky="ew", pady=3)
        e.focus_set()

        ttk.Label(box, text="第二行文字：").grid(row=1, column=0, sticky="w", pady=3)
        ttk.Entry(box, textvariable=self.var_subtitle, width=18).grid(row=1, column=1, sticky="ew", pady=3)

        ttk.Label(box, text="份数：").grid(row=2, column=0, sticky="w", pady=3)
        ttk.Entry(box, textvariable=self.var_copies, width=18).grid(row=2, column=1, sticky="ew", pady=3)

        ttk.Label(box, text="字号：").grid(row=3, column=0, sticky="w", pady=3)
        f = ttk.Frame(box)
        f.grid(row=3, column=1, sticky="ew", pady=3)
        ttk.Scale(f, from_=20, to=60, variable=self.var_fontsize,
                  orient="horizontal", command=lambda v: self.var_fontsize.set(round(float(v)))).pack(side="left", fill="x", expand=True)
        ttk.Label(f, textvariable=self.var_fontsize, width=3).pack(side="left")

        ttk.Label(box, text="上下微调：").grid(row=4, column=0, sticky="w", pady=3)
        f2 = ttk.Frame(box)
        f2.grid(row=4, column=1, sticky="ew", pady=3)
        ttk.Scale(f2, from_=-15, to=15, variable=self.var_voffset,
                  orient="horizontal").pack(side="left", fill="x", expand=True)
        ttk.Label(f2, text="%").pack(side="left")

        ttk.Label(box, text="微调用于对齐不同封皮的印刷位置",
                  foreground="#888", font=("TkDefaultFont", 8)).grid(
            row=5, column=0, columnspan=2, sticky="w")

        # ---- 批量列表 ----
        lbox = ttk.LabelFrame(left, text="批量列表", padding=6)
        lbox.pack(fill="both", expand=True, pady=(10, 0))

        cols = ("cls", "num")
        tree_wrap = ttk.Frame(lbox)
        tree_wrap.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(tree_wrap, columns=cols, show="headings", height=6)
        self.tree.heading("cls", text="班级")
        self.tree.heading("num", text="份数")
        self.tree.column("cls", width=140)
        self.tree.column("num", width=52, anchor="center")
        sb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.tree.bind("<<TreeviewSelect>>", self._on_tree_select)
        self.tree.bind("<Delete>", lambda e: self._del_row())
        self.tree.bind("<Double-1>", lambda e: self._on_tree_select(e))

        bf = ttk.Frame(lbox)
        bf.pack(fill="x", pady=(6, 0))
        ttk.Button(bf, text="加入", width=7, command=self._add_row).pack(side="left", padx=(0, 4))
        ttk.Button(bf, text="更新", width=7, command=self._update_row).pack(side="left", padx=4)
        ttk.Button(bf, text="删除", width=7, command=self._del_row).pack(side="left", padx=4)
        ttk.Button(bf, text="清空", width=7, command=self._clear_rows).pack(side="left", padx=(4, 0))

        # ---- 打印 ----
        pbox = ttk.LabelFrame(left, text="打印", padding=10)
        pbox.pack(fill="x", pady=(10, 0))

        ttk.Label(pbox, text="打印机：").grid(row=0, column=0, sticky="w", pady=3)
        self.cmb_printer = ttk.Combobox(pbox, textvariable=self.var_printer,
                                        width=16, state="readonly")
        self.cmb_printer.grid(row=0, column=1, sticky="ew", pady=3)
        ttk.Button(pbox, text="刷新", width=6, command=self._refresh_printers).grid(
            row=0, column=2, padx=(4, 0))

        ttk.Label(pbox, text="每项份数：").grid(row=1, column=0, sticky="w", pady=3)
        ttk.Spinbox(pbox, from_=1, to=999, textvariable=self.var_pcopies,
                    width=6).grid(row=1, column=1, sticky="w", pady=3)

        self.chk_word = ttk.Checkbutton(
            pbox, text="用 Word 打印（矢量更清晰，需装 Office）",
            variable=self.var_word_print)
        self.chk_word.grid(row=2, column=0, columnspan=3, sticky="w", pady=(4, 0))

        ef = ttk.Frame(pbox)
        ef.grid(row=3, column=0, columnspan=3, sticky="ew", pady=(8, 0))
        ttk.Button(ef, text="导出 PNG", command=self.do_export_png).pack(
            side="left", fill="x", expand=True, padx=(0, 3))
        ttk.Button(ef, text="导出 Word", command=self.do_export_docx).pack(
            side="left", fill="x", expand=True, padx=(3, 0))

        btn = ttk.Frame(left)
        btn.pack(fill="x", pady=(10, 0))
        self.btn_print = ttk.Button(btn, text="打印当前封面", command=self.do_print_current)
        self.btn_print.pack(fill="x", pady=2)
        self.btn_print_all = ttk.Button(btn, text="打印整个列表", command=self.do_print_all)
        self.btn_print_all.pack(fill="x", pady=2)

        self.lbl_status = ttk.Label(left, text="", foreground="#0a0", wraplength=300)
        self.lbl_status.pack(fill="x", pady=(8, 0))

        # ---- 预览 ----
        ttk.Label(right, text="预览效果（与打印结果一致）",
                  font=("TkDefaultFont", 10, "bold")).pack(anchor="w", pady=(0, 6))
        self.canvas = tk.Label(right, bg="#e6e6e6", bd=1, relief="solid")
        self.canvas.pack(fill="both", expand=True)

        if not IS_WIN:
            self.btn_print.configure(state="disabled")
            self.btn_print_all.configure(state="disabled")
            self.cmb_printer["values"] = ["（仅 Windows 可打印）"]
            self.cmb_printer.set("（仅 Windows 可打印）")
            self.chk_word.configure(state="disabled")

        self.bind("<Control-Return>", lambda e: self.do_print_current())

    # ---------------- 预览 ----------------
    def schedule_preview(self):
        if hasattr(self, "_after_id"):
            try:
                self.after_cancel(self._after_id)
            except Exception:
                pass
        self._after_id = self.after(120, self.update_preview)

    def current_layout(self):
        return Layout(self.var_class.get(), self.var_copies.get(),
                      self.var_subtitle.get(), self.var_fontsize.get(),
                      self.var_voffset.get())

    def update_preview(self):
        try:
            lay = self.current_layout()
            img = lay.render(dpi=PREVIEW_DPI)
            max_h = max(200, self.canvas.winfo_height() - 6)
            max_w = max(150, self.canvas.winfo_width() - 6)
            disp = img.copy()
            disp.thumbnail((max_w, max_h), Image.LANCZOS)
            self._photo = make_photo(disp)
            self.canvas.configure(image=self._photo)

            if lay.overflow():
                self.lbl_status.configure(
                    text="内容超出页面，请把字号调小一些。", foreground="#c60")
            elif self.lbl_status.cget("foreground") == "#c60":
                self.lbl_status.configure(text="", foreground="#0a0")
        except Exception as ex:
            self.lbl_status.configure(text=f"预览出错：{ex}", foreground="#c00")

    # ---------------- 列表操作 ----------------
    def _validate(self, quiet=False):
        cls = self.var_class.get().strip()
        num = self.var_copies.get().strip()
        if not cls:
            if not quiet:
                messagebox.showwarning("提示", "请填写班级名称。")
            return None, None
        if not num.isdigit():
            if not quiet:
                messagebox.showwarning("提示", "份数必须是纯数字，例如 47。")
            return None, None
        return cls, num

    def _add_row(self):
        cls, num = self._validate()
        if cls is None:
            return
        for c, n in self._rows:
            if c == cls:
                messagebox.showinfo("提示", f"「{cls}」已经在列表里了。")
                return
        self._rows.append((cls, num))
        self._sync_tree()
        self.lbl_status.configure(text=f"已加入：{cls}（{num} 份）", foreground="#0a0")

    def _update_row(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("提示", "请先在列表里选中一行。")
            return
        cls, num = self._validate()
        if cls is None:
            return
        i = self.tree.index(sel[0])
        self._rows[i] = (cls, num)
        self._sync_tree()
        self.lbl_status.configure(text=f"已更新：{cls}（{num} 份）", foreground="#0a0")

    def _del_row(self):
        sel = self.tree.selection()
        if not sel:
            return
        i = self.tree.index(sel[0])
        del self._rows[i]
        self._sync_tree()

    def _clear_rows(self):
        if not self._rows:
            return
        if messagebox.askyesno("确认", "清空列表？"):
            self._rows = []
            self._sync_tree()

    def _sync_tree(self):
        for iid in self.tree.get_children():
            self.tree.delete(iid)
        for c, n in self._rows:
            self.tree.insert("", "end", values=(c, n))

    def _load_rows(self, rows):
        try:
            for item in rows:
                c, n = item[0], item[1]
                if c and str(n).isdigit():
                    self._rows.append((c, str(n)))
            self._sync_tree()
        except Exception:
            pass

    def _on_tree_select(self, _evt=None):
        sel = self.tree.selection()
        if not sel:
            return
        i = self.tree.index(sel[0])
        if 0 <= i < len(self._rows):
            c, n = self._rows[i]
            self.var_class.set(c)
            self.var_copies.set(n)

    # ---------------- 打印 ----------------
    def do_print_current(self):
        cls, num = self._validate()
        if cls is None:
            return
        self._print_jobs([(cls, num)])

    def do_print_all(self):
        if not self._rows:
            messagebox.showinfo("提示", "列表是空的，先加入几个班级吧。")
            return
        self._print_jobs(list(self._rows))

    def _print_jobs(self, jobs):
        # 先在主线程把所有界面取值取出来，worker 线程不再碰 tkinter
        printer = self.var_printer.get()
        subtitle = self.var_subtitle.get()
        fsize = self.var_fontsize.get()
        voffset = self.var_voffset.get()
        use_word = self.var_word_print.get()
        try:
            pc = int(self.var_pcopies.get())
        except Exception:
            pc = 1

        self.btn_print.configure(state="disabled")
        self.btn_print_all.configure(state="disabled")
        self.lbl_status.configure(text=f"正在打印 {len(jobs)} 项…", foreground="#333")

        def work():
            try:
                for idx, (cls, num) in enumerate(jobs, 1):
                    lay = Layout(cls, num, subtitle, fsize, voffset)
                    if use_word:
                        import tempfile
                        tmp = os.path.join(tempfile.gettempdir(), f"_cover_{cls}.docx")
                        export_docx(lay, tmp)
                        print_via_word(lay, printer, pc, tmp)
                        try:
                            os.remove(tmp)
                        except Exception:
                            pass
                    else:
                        print_layout(lay, printer, pc)
                    self._q.put(("status", f"打印中… {idx}/{len(jobs)}"))
                msg = f"已完成 {len(jobs)} 项打印（{printer or '默认打印机'}，每项 {pc} 份）"
                self._q.put(("done", (msg, True)))
            except Exception as ex:
                self._q.put(("done", (f"打印失败：{ex}", False)))

        threading.Thread(target=work, daemon=True).start()

    def _done(self, msg, ok):
        try:
            self.btn_print.configure(state="normal")
            self.btn_print_all.configure(state="normal")
            self.lbl_status.configure(text=msg, foreground="#0a0" if ok else "#c00")
        except tk.TclError:
            return
        if not ok:
            messagebox.showerror("出错了", msg)

    # ---------------- 导出 ----------------
    def do_export_png(self):
        cls, num = self._validate()
        if cls is None:
            return
        lay = self.current_layout()
        path = filedialog.asksaveasfilename(
            defaultextension=".png",
            initialfile=f"{cls}{lay.subtitle}封面.png",
            filetypes=[("PNG 图片", "*.png")])
        if not path:
            return
        try:
            lay.render(dpi=300).save(path)
            self.lbl_status.configure(text=f"已导出：{os.path.basename(path)}", foreground="#0a0")
        except Exception as ex:
            messagebox.showerror("出错了", str(ex))

    def do_export_docx(self):
        cls, num = self._validate()
        if cls is None:
            return
        lay = self.current_layout()
        path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            initialfile=f"{cls}{lay.subtitle}封面.docx",
            filetypes=[("Word 文档", "*.docx")])
        if not path:
            return
        try:
            export_docx(lay, path)
            self.lbl_status.configure(text=f"已导出：{os.path.basename(path)}", foreground="#0a0")
        except ImportError:
            messagebox.showerror("缺少组件", "未安装 python-docx，无法导出 Word。\n请用「导出 PNG」或直接打印。")
        except Exception as ex:
            messagebox.showerror("出错了", str(ex))

    # ---------------- 杂项 ----------------
    def _refresh_printers(self):
        if not IS_WIN:
            return
        names = list_printers()
        self.cmb_printer["values"] = names
        cur = self.var_printer.get()
        default = get_default_printer()
        if cur in names:
            self.cmb_printer.set(cur)
        elif default and default in names:
            self.cmb_printer.set(default)
        elif names:
            self.cmb_printer.set(names[0])

    def on_close(self):
        save_config({
            "class_name": self.var_class.get(),
            "copies": self.var_copies.get(),
            "subtitle": self.var_subtitle.get(),
            "font_size": self.var_fontsize.get(),
            "voffset": self.var_voffset.get(),
            "printer": self.var_printer.get(),
            "print_copies": self.var_pcopies.get(),
            "word_print": self.var_word_print.get(),
            "rows": [list(r) for r in self._rows],
        })
        self.destroy()


def main():
    if FONT_PATH is None:
        print("警告：未找到中文字体，文字可能显示为方块。")
    App().mainloop()


def cli(class_name, copies, output=None):
    """命令行模式，兼容原来的 make_cover.py 用法。"""
    if not class_name or not str(copies).isdigit():
        print("输入有误：班级名不能为空，份数必须是纯数字。")
        return 1
    lay = Layout(class_name, str(copies))
    if output is None:
        output = f"{class_name}{lay.subtitle}封面.docx"
    export_docx(lay, output)
    print(f"已生成：{output}")
    return 0


if __name__ == "__main__":
    if len(sys.argv) >= 3 and not sys.argv[1].startswith("-"):
        out = sys.argv[3] if len(sys.argv) > 3 else None
        sys.exit(cli(sys.argv[1], sys.argv[2], out))
    elif len(sys.argv) == 2 and sys.argv[1] in ("-h", "--help"):
        print("用法：cover_maker.py [班级名 份数 [输出文件]]")
        print("不带参数则打开图形界面。")
        sys.exit(0)
    else:
        main()
